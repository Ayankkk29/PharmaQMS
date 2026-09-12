import io
import re
import email
from email import policy
import logging
from fastapi import HTTPException, status
from pypdf import PdfReader
import docx

logger = logging.getLogger("pharma_qms.document_parser")

def extract_text_from_file(contents: bytes, filename: str, file_ext: str) -> str:
    """
    Multi-format document text extraction pipeline.
    Parses PDF, DOCX, EML, and TXT files cleanly.
    """
    ext = (file_ext or "").lower()
    extracted_text = ""

    # 1. PDF Parsing (pypdf)
    if ext == ".pdf":
        try:
            pdf_file = io.BytesIO(contents)
            reader = PdfReader(pdf_file)
            text_pages = [page.extract_text() for page in reader.pages if page.extract_text()]
            extracted_text = "\n".join(text_pages).strip()
            logger.info(f"Successfully extracted {len(extracted_text)} chars from PDF file '{filename}' ({len(reader.pages)} pages).")
        except Exception as e:
            logger.error(f"Failed to parse PDF file '{filename}': {e}")
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail="Failed to parse PDF document. Ensure the file is a valid PDF and not corrupt or password-protected."
            )

    # 2. DOCX Parsing (python-docx)
    elif ext in [".docx", ".doc"]:
        try:
            doc = docx.Document(io.BytesIO(contents))
            paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
            
            # Extract table cell texts if present
            table_texts = []
            for table in doc.tables:
                for row in table.rows:
                    row_content = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                    if row_content:
                        table_texts.append(row_content)

            combined = paragraphs + table_texts
            extracted_text = "\n".join(combined).strip()
            logger.info(f"Successfully extracted {len(extracted_text)} chars from DOCX file '{filename}'.")
        except Exception as e:
            logger.error(f"Failed to parse DOCX file '{filename}': {e}")
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail="Failed to parse DOCX document. Ensure the file is a valid Microsoft Word document."
            )

    # 3. EML / Email Parsing (Python email library)
    elif ext in [".eml", ".msg"]:
        try:
            msg = email.message_from_bytes(contents, policy=policy.default)
            subject = msg.get("subject", "No Subject")
            sender = msg.get("from", "Unknown Sender")
            recipient = msg.get("to", "Unknown Recipient")
            date = msg.get("date", "")
            
            body = ""
            if msg.is_multipart():
                for part in msg.walk():
                    content_type = part.get_content_type()
                    disposition = str(part.get("Content-Disposition", ""))
                    if content_type == "text/plain" and "attachment" not in disposition:
                        payload = part.get_payload(decode=True)
                        if payload:
                            body = payload.decode("utf-8", errors="ignore")
                            break
                    elif content_type == "text/html" and not body and "attachment" not in disposition:
                        payload = part.get_payload(decode=True)
                        if payload:
                            html_str = payload.decode("utf-8", errors="ignore")
                            body = re.sub(r"<[^>]+>", " ", html_str)
            else:
                payload = msg.get_payload(decode=True)
                if payload:
                    body = payload.decode("utf-8", errors="ignore")

            extracted_text = f"Email Complaint Intake Report\nFrom: {sender}\nTo: {recipient}\nDate: {date}\nSubject: {subject}\n\nBody Content:\n{body.strip()}"
            logger.info(f"Successfully parsed EML email '{filename}' from sender {sender}.")
        except Exception as e:
            logger.error(f"Failed to parse EML email '{filename}': {e}")
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail="Failed to parse EML file. Ensure the file is a valid MIME email document."
            )

    # 4. TXT / CSV / JSON Plain Text Parsing
    else:
        try:
            extracted_text = contents.decode("utf-8", errors="ignore").strip()
            logger.info(f"Decoded {len(extracted_text)} chars of plain text from file '{filename}'.")
        except Exception as e:
            logger.error(f"Failed to decode text file '{filename}': {e}")
            extracted_text = f"Uploaded document content: {filename}"

    if not extracted_text.strip():
        extracted_text = f"Uploaded complaint file content from {filename}"

    return extracted_text
