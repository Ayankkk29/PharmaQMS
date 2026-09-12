import asyncio
import io
import docx
import httpx
from app.services.document_parser import extract_text_from_file

BASE_URL = "http://127.0.0.1:8000/api/v1"

async def test_parsers_and_chat():
    print("=" * 66)
    print("TESTING MULTI-FORMAT DOCUMENT PARSING & REAL COPILOT CHAT")
    print("=" * 66)

    # 1. Test DOCX Document Parser
    doc = docx.Document()
    doc.add_heading("Customer Complaint - Metformin Hydrochloride API", level=1)
    doc.add_paragraph("Product Name: Metformin Hydrochloride API")
    doc.add_paragraph("Batch Number: B-MET-8841")
    doc.add_paragraph("Defect Description: Discoloration and delayed dissolution observed in drum #4.")
    docx_io = io.BytesIO()
    doc.save(docx_io)
    docx_bytes = docx_io.getvalue()

    extracted_docx = extract_text_from_file(docx_bytes, "test_complaint.docx", ".docx")
    assert "B-MET-8841" in extracted_docx, "DOCX parser failed to extract batch number!"
    print("[PASS] 1. DOCX parser (python-docx) extracted batch and paragraph text successfully.")

    # 2. Test EML Email Parser
    eml_bytes = b"""From: quality@apexpharm.com
To: qms@manufacturer.com
Subject: Quality Defect Report - Paracetamol 500mg Batch PCM240731
Date: Sat, 12 Sep 2026 10:00:00 +0000

Customer reported discoloration in Paracetamol Tablets 500 mg batch PCM240731. 15 strips affected.
"""
    extracted_eml = extract_text_from_file(eml_bytes, "test_email.eml", ".eml")
    assert "PCM240731" in extracted_eml, "EML parser failed to extract batch number!"
    print("[PASS] 2. EML email parser (Python email) extracted headers, subject, and body text successfully.")

    # 3. Test API Endpoint with DOCX upload
    async with httpx.AsyncClient(base_url=BASE_URL, timeout=15.0) as client:
        files = {"file": ("test_complaint.docx", io.BytesIO(docx_bytes), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}
        r = await client.post("/complaints/analyze-file", files=files)
        assert r.status_code == 200, f"Expected 200 for DOCX upload, got {r.status_code}"
        res = r.json()
        ext = res.get("extracted_data", {})
        assert ext.get("batch_number") == "B-MET-8841" or "MET" in str(ext), "LangGraph failed to extract batch from DOCX upload!"
        print("[PASS] 3. End-to-End DOCX upload & LangGraph analysis pipeline verified!")

        # 4. Test Real Copilot Chat Endpoint (POST /api/v1/ai/copilot-chat)
        chat_payload = {
            "query": "What is the recommended containment action for this batch?",
            "complaint_context": res,
            "chat_history": []
        }
        r_chat = await client.post("/ai/copilot-chat", json=chat_payload)
        assert r_chat.status_code == 200, f"Expected 200 for copilot-chat, got {r_chat.status_code}"
        chat_res = r_chat.json()
        assert "reply" in chat_res and len(chat_res["reply"]) > 10, "Copilot chat returned empty response!"
        print(f"[PASS] 4. Real Copilot LLM Chat Endpoint verified! Reply source: {chat_res.get('source')}")
        print(f"       Sample LLM Reply: \"{chat_res['reply'][:120]}...\"")

    print("=" * 66)
    print("DOCUMENT PARSERS & COPILOT CHAT TEST COMPLETED SUCCESSFULLY!")
    print("=" * 66)

if __name__ == "__main__":
    asyncio.run(test_parsers_and_chat())
